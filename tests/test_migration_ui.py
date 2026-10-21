"""User-visible PL/SQL guidance, warning visibility and safe folder publishing."""
from pathlib import Path
import time

import pytest
import yaml
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest


APP = str(Path(__file__).resolve().parents[1] / 'webui' / 'app.py')


def migration_app():
    app=AppTest.from_file(APP,default_timeout=30).run()
    next(r for r in app.radio if r.label=='화면').set_value('SQL 마이그레이션').run()
    assert not app.exception
    return app


def folder_app(tmp_path,source,output=None):
    inputs=tmp_path/'input';inputs.mkdir()
    (inputs/'object.sql').write_text(source)
    mapping=tmp_path/'mapping.yaml'
    mapping.write_text(yaml.safe_dump({'version':'1.0','tables':[{'type':'rename','as_is':'ORDERS','to_be':'TB_ORDER'}],
        'columns':[{'as_is':{'table':'ORDERS','column':'ID'},'to_be':{'table':'TB_ORDER','column':'ORDER_ID'}}]}))
    output=output or tmp_path/'result'
    app=migration_app()
    next(r for r in app.radio if r.label=='실행 방식').set_value('폴더 지정 (지정 폴더로 출력)').run()
    for widget in app.text_input:
        if widget.label.startswith('①'):widget.set_value(str(inputs))
        elif widget.label.startswith('②'):widget.set_value(str(mapping))
        elif widget.label.startswith('③'):widget.set_value(str(output))
    app.run()
    return app,inputs,output


def run_conversion(app):
    next(b for b in app.button if b.label=='▶ 변환 실행').click().run()
    finish_conversion(app)


def finish_conversion(app):
    deadline = time.monotonic() + 30
    while 'mig_job' in app.session_state and time.monotonic() < deadline:
        time.sleep(0.02)
        app.run()
    assert 'mig_job' not in app.session_state, 'conversion did not finish'
    assert not app.exception, [e.message for e in app.exception]


def test_help_explains_examples_and_partial_conversion():
    app=migration_app()
    assert any('AUTO_WARN' in msg.value and 'Conversions' in msg.value for msg in app.info)
    assert any('동적 SQL·모호한 구문이란?' in e.label for e in app.expander)
    code='\n'.join(c.value for c in app.code)
    assert 'EXECUTE IMMEDIATE' in code and 'V_TABLE_NAME' in code
    assert 'SELECT ID' in code and 'SELECT o.ID' in code
    prose='\n'.join(m.value for m in app.markdown)
    assert '고정 문자열이어도' in prose
    assert '컴파일 성공을 의미하지 않습니다' in '\n'.join(c.value for c in app.caption)


def test_upload_result_warnings_are_visible_without_truncated_log():
    app=migration_app()
    warning='warning: object.sql: line 4: dynamic SQL is preserved; inspect EXECUTE IMMEDIATE manually'
    app.session_state['mig_result']={'ok':True,'summary':'Converted 1 statement(s): AUTO=0 AUTO_WARN=1 NEEDS_LLM=0 UNRESOLVED=0 PARSE_FAIL=0',
        'zip':b'zip','log':warning+'\n'+'\n'.join(f'normal line {n}' for n in range(100))}
    app.run()
    assert not app.exception
    assert app.warning and not app.success
    assert any(warning in t.value for t in app.text)
    assert any(e.label=='수동 검토 항목' and e.proto.expanded for e in app.expander)
    assert app.get('download_button')


def test_folder_warning_and_excel_notes(tmp_path):
    source="BEGIN SELECT ID INTO V FROM ORDERS; EXECUTE IMMEDIATE 'SELECT ID FROM ORDERS'; END; /"
    app,inputs,output=folder_app(tmp_path,source)
    run_conversion(app)
    assert app.warning and not app.success and not app.error
    result=(output/'object.sql').read_text()
    assert 'SELECT ORDER_ID INTO V FROM TB_ORDER' in result
    assert "EXECUTE IMMEDIATE 'SELECT ID FROM ORDERS'" in result
    assert (inputs/'object.sql').read_text()==source
    assert any('object.sql: line 1' in t.value for t in app.text)
    workbook=load_workbook(next(output.glob('*.xlsx')),read_only=True)
    rows=list(workbook['Conversions'].values)
    status=rows[0].index('Status');notes=rows[0].index('Notes')
    assert rows[1][status]=='AUTO_WARN'
    assert 'dynamic SQL' in rows[1][notes]
    summary={row[0]:row[1] for row in workbook['Summary'].values if len(row)>1 and row[0]}
    assert summary['Warning-free entries rate']=='0.0% (AUTO only)'
    assert 'Automation rate' not in summary


@pytest.mark.parametrize('relationship',['same','inside','parent'])
def test_related_output_folders_cannot_overwrite_or_reingest_sources(tmp_path,relationship):
    output={'same':tmp_path/'input','inside':tmp_path/'input'/'result','parent':tmp_path}[relationship]
    source='BEGIN SELECT ID INTO V FROM ORDERS; END; /'
    app,inputs,_=folder_app(tmp_path,source,output)
    run_conversion(app)
    assert app.error and '별도 폴더' in app.error[0].value
    assert (inputs/'object.sql').read_text()==source
    assert not list(tmp_path.rglob('*.xlsx'))


def test_failed_batch_does_not_copy_partial_output(tmp_path):
    source='BEGIN SELECT ID INTO V FROM ORDERS; END; /'
    app,inputs,output=folder_app(tmp_path,source)
    (inputs/'z_invalid.sql').write_bytes(b'\xff\xff')
    run_conversion(app)
    assert app.error and not app.success
    assert not list(output.glob('*.sql'))
    assert not list(output.glob('*.xlsx'))
    assert (inputs/'object.sql').read_text()==source


def test_existing_output_symlink_cannot_overwrite_original(tmp_path):
    source='BEGIN SELECT ID INTO V FROM ORDERS; END; /'
    app,inputs,output=folder_app(tmp_path,source)
    output.mkdir()
    (output/'object.sql').symlink_to(inputs/'object.sql')
    run_conversion(app)
    assert app.error
    assert '출력 경로' in app.error[0].value
    assert (inputs/'object.sql').read_text()==source
    assert not list(output.glob('*.xlsx'))


@pytest.mark.parametrize('quoted', [False, True])
def test_xml_only_folder_enables_and_converts_with_pasted_paths(tmp_path, quoted):
    source = '<mapper namespace="test"><select id="get">SELECT ID FROM ORDERS</select></mapper>'
    app, inputs, output = folder_app(tmp_path, 'BEGIN NULL; END; /')
    (inputs / 'object.sql').unlink()
    (inputs / 'nested').mkdir()
    (inputs / 'nested' / 'mapper.xml').write_text(source)
    if quoted:
        for widget in app.text_input:
            if widget.label.startswith(('①', '②', '③')):
                widget.set_value('  "' + widget.value + '"  ')
        app.run()
    assert not next(b for b in app.button if b.label == '▶ 변환 실행').disabled
    run_conversion(app)
    assert not app.error
    converted = (output / 'nested' / 'mapper.xml').read_text()
    assert 'TB_ORDER' in converted and 'ORDER_ID' in converted
    assert (inputs / 'nested' / 'mapper.xml').read_text() == source
    assert list(output.glob('*.xlsx'))


def test_invalid_folder_path_explains_failure_instead_of_silent_disabled_button(tmp_path):
    app, inputs, output = folder_app(tmp_path, 'BEGIN NULL; END; /')
    next(w for w in app.text_input if w.label.startswith('①')).set_value(str(tmp_path / 'missing'))
    app.run()
    assert not next(b for b in app.button if b.label == '▶ 변환 실행').disabled
    run_conversion(app)
    assert any('AS-IS 폴더 경로를 확인하세요' in e.value for e in app.error)
    assert not output.exists()


def test_folder_multiple_updates_are_all_written_and_reported(tmp_path):
    source = ("CREATE OR REPLACE PROCEDURE P AS BEGIN "
              "UPDATE ORDERS SET ID=:ID; UPDATE ORDERS SET ID=2; "
              "UPDATE ORDERS SET ID=3 WHERE ID=2; END; /")
    app, inputs, output = folder_app(tmp_path, source)
    run_conversion(app)
    assert not app.error and app.success
    converted = (output / 'object.sql').read_text()
    assert converted.count('UPDATE TB_ORDER SET ORDER_ID=') == 3
    assert 'WHERE ORDER_ID=2' in converted
    assert (inputs / 'object.sql').read_text() == source
    workbook = load_workbook(next(output.glob('*.xlsx')), read_only=True)
    rows = list(workbook['Conversions'].values)
    assert rows[1][rows[0].index('Status')] == 'AUTO'


def test_folder_result_survives_reruns_and_option_changes(tmp_path):
    app, inputs, output = folder_app(tmp_path, 'BEGIN UPDATE ORDERS SET ID=1; END; /')
    run_conversion(app)
    assert app.success
    app.run()
    assert app.success
    next(c for c in app.checkbox if c.label == '한글 주석 삽입').check().run()
    assert app.success
    assert 'UPDATE TB_ORDER SET ORDER_ID=1' in (output / 'object.sql').read_text()


def test_running_job_survives_option_rerun_without_duplicate_process(tmp_path, monkeypatch):
    import subprocess
    import threading
    app, inputs, output = folder_app(tmp_path, 'BEGIN UPDATE ORDERS SET ID=1; END; /')
    original_run = subprocess.run
    gate = threading.Event()
    started = threading.Event()
    commands = []

    def paused_run(command, *args, **kwargs):
        if isinstance(command, list) and any(str(arg).endswith('main.py') for arg in command):
            commands.append(list(command))
            started.set()
            assert gate.wait(10), 'test did not release conversion'
        return original_run(command, *args, **kwargs)

    monkeypatch.setattr(subprocess, 'run', paused_run)
    try:
        next(b for b in app.button if b.label == '▶ 변환 실행').click().run()
        assert started.wait(2)
        job = app.session_state['mig_job']
        assert not job['future'].done()
        assert next(b for b in app.button if b.label == '▶ 변환 실행').disabled
        next(c for c in app.checkbox if c.label == 'Stage A 검증 건너뛰기').uncheck().run()
        app.run()
        assert not app.exception
        assert app.session_state['mig_job']['future'] is job['future']
        assert any('변환 중' in msg.value for msg in app.info)
        assert len(commands) == 1 and '--no-validate' in commands[0]
    finally:
        gate.set()
    finish_conversion(app)
    assert app.success and (output / 'object.sql').is_file()
    app.run()
    assert app.success and len(commands) == 1


def test_completed_background_failure_is_visible_after_reruns(tmp_path):
    from concurrent.futures import Future
    app, inputs, output = folder_app(tmp_path, 'BEGIN NULL; END; /')
    future = Future()
    future.set_exception(OSError('sample worker failure'))
    app.session_state['mig_job'] = {'mode': 'folder', 'future': future}
    app.run()
    assert not app.exception
    assert any('sample worker failure' in item.value for item in app.code)
    app.run()
    assert app.error and any('sample worker failure' in item.value for item in app.code)


def test_upload_job_and_download_survive_option_reruns():
    from concurrent.futures import Future
    app = migration_app()
    future = Future()
    app.session_state['mig_job'] = {'mode': 'upload', 'future': future}
    app.run()
    next(c for c in app.checkbox if c.label == '한글 주석 삽입').check().run()
    assert app.session_state['mig_job']['future'] is future
    assert any('변환 중' in msg.value for msg in app.info)
    future.set_result({'ok': True, 'summary': 'Converted 1 statement(s): AUTO=1 AUTO_WARN=0',
                       'zip': b'zip-result', 'log': 'conversion complete'})
    app.run()
    assert app.success and app.get('download_button')
    app.run()
    assert app.success and app.get('download_button')
    assert app.session_state['mig_result']['zip'] == b'zip-result'
